from openpyxl import Workbook
from datetime import datetime

# Создаем файл plavka.xlsx
def create_plavka_file():
    wb = Workbook()
    ws = wb.active
    
    # Заголовки
    headers = ['№ п/п', 'Учетный номер', 'Дата плавки', 'Масса, кг', 'Температура, C',
               'Форма', 'Плавильщик', 'Литейщик', 'Мастер', 'Заказчик', 'Наименование отливки', 'Тип эксперимента']
    
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    
    # Тестовые данные
    sample_data = [
        [1, '123/25', datetime(2025, 5, 10), 25.5, 1450, 'Форма-1', 'Иванов', 'Петров', 'Сидоров', 'ООО Металл', 'Корпус насоса', 'Стандартный'],
        [2, '124/25', datetime(2025, 5, 11), 30.2, 1460, 'Форма-2', 'Иванов', 'Смирнов', 'Сидоров', 'ООО Металл', 'Лопатка турбины', 'Специальный'],
        [3, '125/25', datetime(2025, 5, 12), 15.7, 1440, 'Форма-3', 'Смирнов', 'Петров', 'Кузнецов', 'ЗАО ТехМет', 'Крышка редуктора', 'Новый сплав'],
        [4, '126/25', datetime(2025, 5, 13), 22.0, 1455, 'Форма-1', 'Иванов', 'Кузнецов', 'Сидоров', 'ОАО Промдеталь', 'Вал двигателя', 'Стандартный'],
        [5, '127/25', datetime(2025, 5, 14), 18.3, 1450, 'Форма-2', 'Смирнов', 'Петров', 'Кузнецов', 'ООО Металл', 'Фланец', 'Оптимизация']
    ]
    
    for row_idx, row_data in enumerate(sample_data, start=2):
        for col_idx, cell_value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=cell_value)
    
    # Настройка формата для дат
    for row in range(2, len(sample_data) + 2):
        ws.cell(row=row, column=3).number_format = 'DD.MM.YYYY'
    
    wb.save('plavka.xlsx')
    print("Файл plavka.xlsx создан успешно!")

# Создаем пустой файл marshrutka.xlsx
def create_marshrutka_file():
    wb = Workbook()
    ws = wb.active
    ws.title = "Records"
    
    # Заголовки
    headers = ['Дата сборки', 'Специалист сборки', 'Количество', 
              'Дата выставления', 'Время выставления', 'Специалист контроля',
              'Учетный номер', 'Наименование отливки', 'Тип эксперимента',
              'Дата болгарки', 'Специалист болгарки', 'Специалист термообработки',
              'Специалист дробеметной обработки', 'Специалист зачистки короны',
              'Специалист зачистки лапы', 'Специалист зачистки питателя', 'Примечание']
    
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    
    # Задаем форматы для столбцов с датами и временем
    for col in [1, 4, 10]:  # Колонки с датами
        ws.column_dimensions[chr(64 + col)].number_format = 'DD.MM.YYYY'
    
    # Колонка с временем
    ws.column_dimensions[chr(64 + 5)].number_format = 'HH:MM'
    
    wb.save('marshrutka.xlsx')
    print("Файл marshrutka.xlsx создан успешно!")

if __name__ == "__main__":
    create_plavka_file()
    create_marshrutka_file() 